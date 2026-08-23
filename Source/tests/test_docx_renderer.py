import io
import unittest

from docx import Document

from certificate_engine.docx_renderer import (
    generate_docx_certificate,
    load_docx_template,
)
from certificate_engine.exceptions import DOCXValidationError

from .helpers import configuration, make_docx


class DocxRendererTests(unittest.TestCase):
    def setUp(self):
        self.configuration = configuration()

    def _document_text(self, raw):
        document = Document(io.BytesIO(raw))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def test_loads_valid_template_and_discovers_placeholders(self):
        template = load_docx_template(make_docx())

        self.assertEqual(
            template.placeholders,
            frozenset({"recipient_name", "certificate_number"}),
        )

    def test_renders_configured_values(self):
        output = generate_docx_certificate(
            make_docx(),
            self.configuration,
            {
                "recipient_name": "Amina & Naledi",
                "certificate_number": "CERT-001",
            },
        )

        text = self._document_text(output)
        self.assertIn("Certificate for Amina & Naledi", text)
        self.assertIn("Number CERT-001", text)

    def test_rejects_empty_template(self):
        with self.assertRaisesRegex(DOCXValidationError, "empty"):
            load_docx_template(b"")

    def test_rejects_malformed_template(self):
        with self.assertRaisesRegex(DOCXValidationError, "malformed"):
            load_docx_template(b"not-a-docx")

    def test_rejects_template_missing_configured_placeholder(self):
        with self.assertRaisesRegex(
            DOCXValidationError, "certificate_number"
        ):
            generate_docx_certificate(
                make_docx("Certificate for {{ recipient_name }}"),
                self.configuration,
                {
                    "recipient_name": "Amina",
                    "certificate_number": "CERT-001",
                },
            )
